from io import BytesIO
from uuid import uuid4

from docx import Document
from starlette.testclient import TestClient

from app.main import app
from app.services.lesson_note_export_service import (
    export_lesson_note_docx,
    export_lesson_note_pdf,
    lesson_note_filename,
)


def test_lesson_note_pdf_renders_markdown_and_readable_inline_math():
    from pypdf import PdfReader

    payload = export_lesson_note_pdf(
        title="QA Lesson Note",
        subject="Mathematics",
        class_level="SSS 3",
        term="First Term",
        week_number=1,
        school_name="QA School",
        content=(
            "**LESSON NOTE: MATHEMATICS**\n"
            "---\n"
            "## Objectives\n"
            "- **Recall:** apply the law $\\log_{10}(100)=2$.\n"
            "- Compare \\(x^2 + \\frac{1}{2}\\).\n"
        ),
    )

    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(payload)).pages)

    assert "LESSON NOTE: MATHEMATICS" in text
    assert "Recall:" in text
    assert "log10(100)=2" in text
    assert "x2 + 1/2" in text
    assert "**" not in text
    assert "\\log" not in text
    assert "---" not in text


def test_lesson_note_pdf_embeds_uploaded_school_logo_and_watermark():
    from pypdf import PdfReader

    payload = export_lesson_note_pdf(
        title="Quadratic Equations",
        subject="Mathematics",
        class_level="SSS 3",
        term="First Term",
        week_number=2,
        school_name="QA School",
        school_logo="app/assets/img/skuphase_logo.png",
        content=r"Apply $x^2 + \frac{1}{2}$ and log_a(xy).",
    )
    page = PdfReader(BytesIO(payload)).pages[0]
    text = page.extract_text()

    assert page.images
    assert text.count("QA School") >= 2  # document heading plus faint diagonal watermark
    assert "x2" in text and "1/2" in text and "loga(xy)" in text


def test_lesson_note_filename_uses_safe_title_and_short_unique_id():
    assert lesson_note_filename("Primary 6: Quadratic Equations!", "12345678-90ab", "pdf") == "primary-6-quadratic-equations-12345678.pdf"


def test_school_webp_logo_is_normalized_for_pdf_and_word_export(tmp_path):
    from PIL import Image
    from app.services.branding import get_logo_bytes

    path = tmp_path / "school-logo.webp"
    Image.new("RGB", (48, 24), (24, 100, 70)).save(path, format="WEBP")

    normalized = get_logo_bytes(str(path))

    assert normalized.startswith(b"\x89PNG\r\n\x1a\n")


def test_lesson_note_docx_preserves_metadata_and_structured_content():
    payload = export_lesson_note_docx(
        title="QA Lesson Note",
        subject="Mathematics",
        class_level="Primary 6",
        term="Second Term",
        week_number=4,
        school_name="QA School",
        teacher_name="QA Teacher",
        content=(
            "# Objectives\n"
            "- Identify equal fractions\n"
            "- Compare \\(a^2+b^2\\)\n"
            "1. Explain the method\n"
            "| Item | Value |\n"
            "| --- | --- |\n"
            "| Fractions | 1/2 |\n"
            "---\n"
        ),
        school_logo="app/assets/img/skuphase_logo.png",
    )

    document = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )

    for expected in (
        "QA School",
        "QA Lesson Note",
        "Mathematics",
        "Primary 6",
        "Second Term",
        "Week 4",
        "QA Teacher",
        "Identify equal fractions",
        r"Compare \(a^2+b^2\)",
        "Explain the method",
        "Fractions",
    ):
        assert expected in text

    assert len(document.tables) == 1
    assert len(document.inline_shapes) == 1
    assert "Official school copy" in document.sections[0].footer.paragraphs[0].text
    assert "QA SCHOOL" in document.sections[0].header.paragraphs[0].text
    assert "---" not in text
    assert any(paragraph.style.name == "List Bullet" for paragraph in document.paragraphs)
    assert any(paragraph.style.name == "List Number" for paragraph in document.paragraphs)


def test_lesson_note_docx_marks_personal_branding_as_personal_copy():
    payload = export_lesson_note_docx(
        title="Custom lesson",
        subject="Science",
        class_level="Primary 5",
        term="First Term",
        week_number=1,
        school_name="Tutor Ada",
        content="A short note.",
        is_personal_workspace=True,
    )

    document = Document(BytesIO(payload))
    footer = document.sections[0].footer.paragraphs[0].text
    assert "Tutor Ada" in footer
    assert "Personal teaching copy" in footer
    assert "Official school copy" not in footer


def test_lesson_note_docx_frontend_proxy_requires_session_and_preserves_download(monkeypatch):
    client = TestClient(app, follow_redirects=False)
    lesson_id = uuid4()
    response = client.get(f"/app/teaching/{lesson_id}/lesson-note/export.docx")
    assert response.status_code == 303

    class StubResponse:
        def __init__(self, status_code, payload=None, content=b"", headers=None):
            self.status_code = status_code
            self._payload = payload or {}
            self.content = content
            self.headers = headers or {}
            self.is_success = 200 <= status_code < 300

        def json(self):
            return self._payload

    class StubApi:
        def __init__(self, response):
            self.response = response
            self.calls = []

        async def __call__(self, req, method, path, json=None, params=None):
            self.calls.append((method, path))
            return self.response

    token_stub = StubApi(StubResponse(200, {
        "access_token": "at",
        "refresh_token": "rt",
        "token_type": "bearer",
        "expires_in": 3600,
        "user": {
            "user_id": "u1",
            "full_name": "QA Teacher",
            "email": "qa@example.test",
            "role": "teacher",
            "account_type": "school_staff",
            "is_active": True,
            "is_verified": True,
        },
    }))
    docx_bytes = b"PK\x03\x04test-docx"
    download_stub = StubApi(StubResponse(
        200,
        content=docx_bytes,
        headers={"content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ))
    monkeypatch.setattr("app.frontend.routes.auth.call_api", token_stub)
    monkeypatch.setattr("app.frontend.routes.teaching.call_api", download_stub)

    login_response = client.post("/login", data={"email": "qa@example.test", "password": "pw"})
    assert login_response.status_code in (200, 303)
    response = client.get(f"/app/teaching/{lesson_id}/lesson-note/export.docx")

    assert response.status_code == 200
    assert response.content == docx_bytes
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.headers["content-disposition"] == f'attachment; filename="lesson-note-{str(lesson_id)[:8]}.docx"'
    assert download_stub.calls == [("GET", f"/lesson-plans/{lesson_id}/lesson-note/export.docx")]


def test_lesson_note_pdf_frontend_proxy_keeps_title_based_filename(monkeypatch):
    client = TestClient(app, follow_redirects=False)
    lesson_id = uuid4()

    class StubResponse:
        def __init__(self, status_code, payload=None, content=b"", headers=None):
            self.status_code = status_code
            self._payload = payload or {}
            self.content = content
            self.headers = headers or {}
            self.is_success = 200 <= status_code < 300

        def json(self):
            return self._payload

    class StubApi:
        def __init__(self, response):
            self.response = response

        async def __call__(self, req, method, path, json=None, params=None):
            return self.response

    token_stub = StubApi(StubResponse(200, {
        "access_token": "at", "refresh_token": "rt", "token_type": "bearer", "expires_in": 3600,
        "user": {"user_id": "u1", "full_name": "QA Teacher", "email": "qa@example.test", "role": "teacher", "account_type": "school_staff", "is_active": True, "is_verified": True},
    }))
    pdf_stub = StubApi(StubResponse(
        200,
        content=b"%PDF-test",
        headers={"content-type": "application/pdf", "content-disposition": 'attachment; filename="quadratic-equations-12345678.pdf"'},
    ))
    monkeypatch.setattr("app.frontend.routes.auth.call_api", token_stub)
    monkeypatch.setattr("app.frontend.routes.teaching.call_api", pdf_stub)
    assert client.post("/login", data={"email": "qa@example.test", "password": "pw"}).status_code == 303

    response = client.get(f"/app/teaching/{lesson_id}/lesson-note/export.pdf")

    assert response.status_code == 200
    assert response.content == b"%PDF-test"
    assert response.headers["content-disposition"] == 'attachment; filename="quadratic-equations-12345678.pdf"'
