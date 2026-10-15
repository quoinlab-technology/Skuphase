"""PDF rendering for saved teacher lesson notes."""

from __future__ import annotations

import re
from io import BytesIO
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from app.services.branding import get_logo_bytes, logo_flowable


_INLINE_TOKEN_RE = re.compile(
    r"(\$[^$]+\$|\\\(.+?\\\)|\([^()]+\)/\([^()]+\)|\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))"
)


def _readable_inline_formula(source: str) -> str:
    """Render common inline TeX with typographic sub/superscripts and symbols."""
    value = source.strip().strip("$")
    if value.startswith(r"\(") and value.endswith(r"\)"):
        value = value[2:-2]
    value = re.sub(r"^\(([^()]+)\)/\(([^()]+)\)$", r"\\frac{\1}{\2}", value)
    symbols = {
        "times": "×", "div": "÷", "pm": "±", "cdot": "·", "leq": "≤",
        "le": "≤", "geq": "≥", "ge": "≥", "neq": "≠", "rightarrow": "→",
        "to": "→", "pi": "π", "theta": "θ", "Delta": "Δ", "lambda": "λ",
        "infty": "∞", "alpha": "α", "beta": "β", "gamma": "γ", "mu": "μ",
        "sigma": "σ", "omega": "ω", "sqrt": "√", "log": "log", "ln": "ln",
        "sin": "sin", "cos": "cos", "tan": "tan",
    }
    symbols.update({
        "times": "\u00d7", "div": "\u00f7", "pm": "\u00b1", "cdot": "\u00b7",
        "leq": "\u2264", "le": "\u2264", "geq": "\u2265", "ge": "\u2265",
        "neq": "\u2260", "rightarrow": "\u2192", "to": "\u2192", "pi": "\u03c0",
        "theta": "\u03b8", "Delta": "\u0394", "lambda": "\u03bb", "infty": "\u221e",
        "alpha": "\u03b1", "beta": "\u03b2", "gamma": "\u03b3", "mu": "\u03bc",
        "sigma": "\u03c3", "omega": "\u03c9", "sqrt": "\u221a",
    })

    def group_at(text: str, start: int) -> tuple[str, int]:
        if start >= len(text) or text[start] != "{":
            return text[start:start + 1], start + 1
        depth = 1
        cursor = start + 1
        while cursor < len(text) and depth:
            if text[cursor] == "{": depth += 1
            elif text[cursor] == "}": depth -= 1
            cursor += 1
        return text[start + 1:cursor - 1], cursor

    def parse(text: str) -> str:
        parts = []
        i = 0
        while i < len(text):
            if text[i] == "\\":
                match = re.match(r"\\([A-Za-z]+)", text[i:])
                if match:
                    command = match.group(1)
                    i += len(match.group(0))
                    if command in {"left", "right"}:
                        continue
                    if command in {"text", "mathrm", "mathbf", "operatorname"} and i < len(text) and text[i] == "{":
                        group, i = group_at(text, i)
                        parts.append(parse(group))
                    elif command == "frac":
                        numerator, i = group_at(text, i)
                        denominator, i = group_at(text, i)
                        parts.append(f"<super>{parse(numerator)}</super>/<sub>{parse(denominator)}</sub>")
                    else:
                        parts.append(escape(symbols.get(command, command)))
                    continue
                if i + 1 < len(text) and text[i + 1] in "{}%&#_$":
                    parts.append(escape(text[i + 1]))
                    i += 2
                    continue
            if text[i] in "_^":
                tag = "sub" if text[i] == "_" else "super"
                group, next_i = group_at(text, i + 1)
                if next_i > i + 1:
                    parts.append(f"<{tag}>{parse(group)}</{tag}>")
                    i = next_i
                    continue
            if text[i] == "{":
                group, i = group_at(text, i)
                parts.append(parse(group))
                continue
            parts.append(escape(text[i]))
            i += 1
        return "".join(parts)

    return re.sub(r"\\[,;!]", " ", parse(value))


def _reportlab_inline(text: str) -> str:
    """Escape untrusted content while supporting safe Markdown emphasis/code."""
    rendered: list[str] = []
    for token in _INLINE_TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("$") or token.startswith(r"\(") or re.fullmatch(r"\([^()]+\)/\([^()]+\)", token):
            rendered.append(_readable_inline_formula(token))
        elif token.startswith("**") and token.endswith("**"):
            rendered.append(f"<b>{escape(token[2:-2])}</b>")
        elif token.startswith("*") and token.endswith("*"):
            rendered.append(f"<i>{escape(token[1:-1])}</i>")
        elif token.startswith("`") and token.endswith("`"):
            rendered.append(f'<font name="Courier">{escape(token[1:-1])}</font>')
        else:
            plain = escape(token)
            plain = re.sub(r"(?<![A-Za-z0-9])([A-Za-z]+)_([A-Za-z0-9]+)", r"\1<sub>\2</sub>", plain)
            plain = re.sub(r"([A-Za-z0-9)])\^([A-Za-z0-9]+)", r"\1<super>\2</super>", plain)
            rendered.append(plain)
    return "".join(rendered)


def export_lesson_note_pdf(
    *,
    title: str,
    subject: str,
    class_level: str,
    term: str,
    week_number: int,
    content: str,
    school_name: str | None = None,
    school_logo: str | None = None,
) -> bytes:
    """Render a saved lesson note to a printable, safely escaped PDF."""
    _register_lesson_font()
    stream = BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title[:200],
        author=(school_name or "SkuPhase")[:200],
    )
    styles = getSampleStyleSheet()
    school_style = ParagraphStyle(
        "LessonNoteSchool", parent=styles["Normal"], alignment=TA_CENTER,
        fontSize=10, leading=13, textColor=colors.HexColor("#53615B"), spaceAfter=2 * mm,
    )
    title_style = ParagraphStyle(
        "LessonNoteTitle", parent=styles["Title"], alignment=TA_CENTER,
        fontSize=17, leading=21, textColor=colors.HexColor("#064E3B"), spaceAfter=3 * mm,
    )
    meta_style = ParagraphStyle(
        "LessonNoteMeta", parent=styles["Normal"], alignment=TA_CENTER,
        fontSize=9, leading=13, textColor=colors.HexColor("#53615B"), spaceAfter=7 * mm,
    )
    body_style = ParagraphStyle(
        "LessonNoteBody", parent=styles["BodyText"], fontName="SkuVera", fontSize=10, leading=15,
        spaceAfter=2.5 * mm, wordWrap="CJK",
    )
    heading_style = ParagraphStyle(
        "LessonNoteSection", parent=styles["Heading2"], fontName="SkuVera-Bold", fontSize=12, leading=16,
        textColor=colors.HexColor("#064E3B"), spaceBefore=3 * mm, spaceAfter=1.5 * mm,
    )

    story = []
    logo = logo_flowable(school_logo, width=34 * mm, height=17 * mm)
    if logo is not None:
        logo.hAlign = "CENTER"
        story.extend((logo, Spacer(1, 2 * mm)))
    if school_name:
        story.append(Paragraph(escape(school_name), school_style))
    story.append(Paragraph(escape(title), title_style))
    meta = " | ".join(
        part for part in (
            f"{subject} | {class_level}" if subject or class_level else "",
            term,
            f"Week {week_number}",
        ) if part
    )
    story.append(Paragraph(escape(meta), meta_style))

    normalized_content = (content or "").replace("\u2022 ", "- ")
    for raw_line in normalized_content.splitlines():
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 2 * mm))
            continue
        if re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", line):
            story.append(Spacer(1, 1.5 * mm))
            continue
        heading = re.sub(r"^#{1,3}\s*", "", line).rstrip(":").strip()
        if line.startswith("#") or heading.casefold() in {
            "objectives", "learning objectives", "starter", "introduction",
            "explanation", "guided practice", "assessment", "differentiation",
            "teaching materials", "resources", "conclusion",
        }:
            story.append(Paragraph(_reportlab_inline(heading), heading_style))
        elif line.startswith(("- ", "* ", "• ", "â€¢ ")):
            story.append(Paragraph(_reportlab_inline(line[2:]), body_style, bulletText="-"))
        elif re.match(r"^\d+[.)]\s+", line):
            story.append(Paragraph(_reportlab_inline(line), body_style))
        else:
            story.append(Paragraph(_reportlab_inline(line), body_style))

    def page_furniture(canvas, doc):
        canvas.saveState()
        if school_name:
            canvas.setFillColor(colors.HexColor("#DCE7E0"))
            canvas.setFont("SkuVera-Bold", 23)
            canvas.translate(A4[0] / 2, A4[1] / 2)
            canvas.rotate(34)
            canvas.drawCentredString(0, 0, school_name[:42])
            canvas.restoreState()
            canvas.saveState()
        canvas.setFillColor(colors.HexColor("#64736C"))
        canvas.setFont("SkuVera", 8)
        canvas.drawCentredString(A4[0] / 2, 10 * mm, f"{school_name or 'SkuPhase'} | {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=page_furniture, onLaterPages=page_furniture)
    return stream.getvalue()


def export_lesson_note_docx(
    *,
    title: str,
    subject: str,
    class_level: str,
    term: str,
    week_number: int,
    content: str,
    school_name: str | None = None,
    teacher_name: str | None = None,
    school_logo: str | None = None,
    is_personal_workspace: bool = False,
) -> bytes:
    """Build an editable Word lesson note from the saved Markdown-like text.

    Headings, lists, and simple pipe tables become native Word structures.
    Inline math remains as literal TeX source (for example ``$a^2+b^2$``),
    rather than being silently discarded or misrepresented as rendered math.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)

    if school_logo:
        try:
            image_bytes = get_logo_bytes(school_logo)
            if image_bytes:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(BytesIO(image_bytes), width=Inches(1.0))
        except Exception as exc:
            import logging
            logging.getLogger(__name__).warning("Word export school logo unavailable (%s)", type(exc).__name__)
    if school_name:
        school = document.add_paragraph()
        school.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = school.add_run(school_name.strip())
        run.bold = True
        run.font.color.rgb = RGBColor(6, 78, 59)
        run.font.size = Pt(11)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = heading.add_run(title.strip() or "Lesson Note")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(6, 78, 59)
    title_run.font.size = Pt(17)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = " | ".join(part for part in (
        f"Subject: {subject}" if subject else "",
        f"Class: {class_level}" if class_level else "",
        f"Term: {term}" if term else "",
        f"Week {week_number}" if week_number else "",
        f"Teacher: {teacher_name}" if teacher_name else "",
    ) if part)
    metadata.add_run(meta).font.color.rgb = RGBColor(83, 97, 91)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_label = "Personal teaching copy" if is_personal_workspace else "Official school copy"
    footer_run = footer.add_run(f"{school_name or 'SkuPhase'} · {footer_label}")
    footer_run.font.color.rgb = RGBColor(175, 190, 181)
    footer_run.font.size = Pt(8)

    watermark = section.header.paragraphs[0]
    watermark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    watermark_run = watermark.add_run((school_name or "SkuPhase").strip().upper())
    watermark_run.bold = True
    watermark_run.font.color.rgb = RGBColor(224, 234, 228)
    watermark_run.font.size = Pt(16)

    known_headings = {
        "objectives", "learning objectives", "starter", "introduction",
        "explanation", "guided practice", "assessment", "differentiation",
        "teaching materials", "resources", "conclusion",
    }

    def add_inline(paragraph, text: str) -> None:
        # Preserve TeX literally; translate only simple Markdown emphasis/code.
        tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`|\$[^$]+\$)", text)
        for token in tokens:
            if not token:
                continue
            if token.startswith("**") and token.endswith("**"):
                paragraph.add_run(token[2:-2]).bold = True
            elif token.startswith("`") and token.endswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
            else:
                paragraph.add_run(token)

    def table_cells(line: str) -> list[str]:
        return [cell.strip() for cell in line.strip().strip("|").split("|")]

    def is_table_separator(line: str) -> bool:
        cells = table_cells(line)
        return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)

    lines = (content or "").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", line):
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1].strip()):
            rows = [table_cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(table_cells(lines[index]))
                index += 1
            column_count = max(len(row) for row in rows)
            table = document.add_table(rows=1, cols=column_count)
            table.style = "Table Grid"
            table.autofit = True
            for row_index, values in enumerate(rows):
                cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
                for col_index in range(column_count):
                    cell = cells[col_index]
                    cell.text = values[col_index] if col_index < len(values) else ""
                    if row_index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            document.add_paragraph()
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            paragraph = document.add_heading(heading_match.group(2).strip(), level=level)
        elif line.rstrip(":").casefold() in known_headings:
            paragraph = document.add_heading(line.rstrip(":").strip(), level=2)
        elif re.match(r"^(?:[-*]|•)\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^(?:[-*]|•)\s+", "", line))
        elif re.match(r"^\d+[.)]\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+[.)]\s+", "", line))
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
        index += 1

    if not (content or "").strip():
        document.add_paragraph("No lesson-note content was provided.")

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _register_lesson_font() -> None:
    """Register ReportLab's bundled Vera fonts for reliable math symbols."""
    import reportlab
    from pathlib import Path

    font_dir = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("SkuVera", "Vera.ttf"), ("SkuVera-Bold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            font_path = font_dir / filename
            if font_path.is_file():
                pdfmetrics.registerFont(TTFont(name, str(font_path)))


def lesson_note_filename(title: str, plan_id: str, extension: str) -> str:
    """Make a portable, collision-resistant export filename from its title."""
    slug = re.sub(r"[^a-z0-9]+", "-", (title or "lesson-note").casefold()).strip("-")[:64]
    return f"{slug or 'lesson-note'}-{plan_id[:8]}.{extension.lstrip('.')}"
